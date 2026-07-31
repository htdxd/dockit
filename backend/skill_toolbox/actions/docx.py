from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def _require_text(spec: dict[str, Any], key: str) -> str:
    value = spec.get(key)
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"spec.{key} must be a non-empty string")
    return value.strip()


def _set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _add_formula(document: Document, formula: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_paragraph = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = formula
    run.append(text)
    math.append(run)
    math_paragraph.append(math)
    paragraph._p.append(math_paragraph)


def _add_code_block(document: Document, code: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F4F6")
    properties.append(shading)
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def build_simple_docx(spec_path: Path, output_path: Path) -> None:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    title = _require_text(spec, "title")
    summary = _require_text(spec, "summary")
    formula = _require_text(spec, "formula")
    code = _require_text(spec, "code")
    table_spec = spec.get("table")
    if not isinstance(table_spec, dict):
        raise TypeError("spec.table must be an object")
    headers = table_spec.get("headers")
    rows = table_spec.get("rows")
    if not isinstance(headers, list) or not headers:
        raise ValueError("spec.table.headers must be a non-empty list")
    if not isinstance(rows, list) or not rows:
        raise ValueError("spec.table.rows must be a non-empty list")
    if any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
        raise ValueError("Each table row must match the header count")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    title_paragraph = document.add_heading(title, level=0)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(summary)

    document.add_heading("结构化表格", level=1)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        _set_cell_shading(cell, "DCE6F1")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)

    document.add_heading("公式", level=1)
    _add_formula(document, formula)
    document.add_heading("代码示例", level=1)
    _add_code_block(document, code)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
