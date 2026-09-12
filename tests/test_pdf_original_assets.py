"""PDF 材料生产入口保留原生图片，不依赖额外上传原图。"""

import io
import json
import zipfile

import fitz
import pytest
from docx import Document
from PIL import Image

from skill_toolbox.materials import MaterialService, material_id_dir, _docx_assets
from skill_toolbox.tools.mineru import MineruService
from skill_toolbox.tools.workspace import sha256_file


def test_generated_docx_photo_uses_opc_type_when_read_again(tmp_path):
    photo = io.BytesIO()
    Image.new("RGB", (20, 21), "blue").save(photo, format="PNG")
    path = tmp_path / "resume.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/media/image1.jpeg", photo.getvalue())
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/media/image1.jpeg" ContentType="image/png"/></Types>')
    with zipfile.ZipFile(path) as archive:
        assets, _ = _docx_assets(archive, tmp_path, "resume.docx")
    assert assets[0].mime_type == "image/png"
    assert (tmp_path / assets[0].path).read_bytes() == photo.getvalue()


@pytest.mark.parametrize("reconstructed", [True, False])
def test_pdf_material_keeps_original_image_bytes_and_pixels(tmp_path, monkeypatch, reconstructed):
    original = io.BytesIO()
    Image.new("RGB", (203, 204), (32, 64, 128)).save(original, format="JPEG")
    original_bytes = original.getvalue()
    source = tmp_path / "resume.pdf"
    with fitz.open() as pdf:
        first = pdf.new_page()
        xref = first.insert_image(fitz.Rect(20, 20, 120, 125), stream=original_bytes)
        second = pdf.new_page()
        second.insert_image(fitz.Rect(20, 20, 120, 125), xref=xref)
        pdf.save(source)
    source_hash = sha256_file(source)

    def convert(_self, pdf_path, output, options):
        # Simulate MinerU's reconstructed image while exercising real DOCX/IR parsing.
        photo = io.BytesIO(original_bytes)
        if reconstructed:
            photo = io.BytesIO()
            Image.new("RGB", (206, 210), "gray").save(photo, format="JPEG")
            photo.seek(0)
        document = Document()
        document.add_paragraph("工作经历保持原有解析结果")
        document.add_picture(photo)
        document.save(output)
        return {"status": "miss", "cache_key": "local-test"}

    monkeypatch.setattr(MineruService, "convert", convert)
    workspace = tmp_path / "workspace"
    service = MaterialService(workspace, cache_root=tmp_path / "cache")
    ir = service.prepare_material(source)
    originals = [a for a in ir.assets if a.source_locator.startswith("pdf:embedded-image:")]
    assert len(originals) == 1  # One image reused on two pages, one asset.
    asset = originals[0]
    assert (asset.width, asset.height) == (203, 204)
    assert (workspace / asset.path).read_bytes() == original_bytes
    assert "未裁剪或缩放" in asset.caption
    assert "整页扫描图" in asset.caption
    assert sha256_file(source) == source_hash
    assert any("工作经历" in block.text for block in ir.blocks)
    assert len(ir.assets) == (2 if reconstructed else 1)
    if not reconstructed:
        assert any(asset.id in block.asset_ids for block in ir.blocks)
    stored = json.loads((material_id_dir(workspace, ir.material_id) / "document.json").read_text(encoding="utf-8"))
    assert any(a["id"] == asset.id and a["source_locator"] == asset.source_locator for a in stored["assets"])
    assert service.catalog().asset_path(asset.id) == asset.path
