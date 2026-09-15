"""阶段 2 强制测试：最小 DocumentIR 与隔离存储。

覆盖实施计划 §15.2/§15.3 关键回归：
- 同名文件（不同目录）、同 stem 不同扩展名互不覆盖
- Markdown 相对资源物化 + 缺失引用 warning + `..` 越界拒绝
- 图片块保留资产位置引用；表格/公式/代码不因 Markdown 投影消失
- 同 hash 只解析一次（manifest 记录多个原始名称）
- DocumentIR 由 Pydantic 校验（非法 schema/bbox/ContentPlan 拒绝）
- content.md 兼容投影单向生成
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pydantic import ValidationError
from skill_toolbox.material_models import (
    SCHEMA_VERSION,
    Block,
    ContentPlan,
    DocumentIR,
    Selection,
)
from skill_toolbox.materials import (
    MaterialError,
    MaterialService,
    material_id_dir,
    material_id_for,
    project_content_md,
    safe_stem,
)


def _write(path: Path, content: str | bytes) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(content, bytes):
        path.write_bytes(content)
    else:
        path.write_text(content, encoding="utf-8")
    return path


# ---------------- 数据模型单测 ----------------

def test_document_ir_defaults_and_validation() -> None:
    ir = DocumentIR(
        material_id="abc.pdf",
        original_name="报告.pdf",
        source_format="pdf",
        sha256="x" * 64,
        blocks=[
            Block(id="b0", type="heading", order=0, text="标题", level=1, page=1),
            Block(
                id="b1",
                type="image",
                order=1,
                page=1,
                asset_ids=["a0"],
                bbox=[0.1, 0.1, 0.5, 0.4],
            ),
        ],
        assets=[],
    )
    assert ir.schema_version == SCHEMA_VERSION
    assert ir.block_by_id("b1") is not None
    assert ir.block_by_id("missing") is None
    assert ir.asset_ids_in_use() == {"a0"}


def test_document_ir_rejects_bad_bbox_and_version() -> None:
    with pytest.raises(ValidationError):
        Block(id="b", type="image", order=0, bbox=[0.1, 0.1, 0.5])  # 3 元素
    with pytest.raises(ValidationError):
        Block(id="b", type="image", order=0, bbox=[0.1, 0.1, 0.5, 2.0])  # 越界
    with pytest.raises(ValidationError):
        DocumentIR(
            material_id="a",
            original_name="a.md",
            source_format="md",
            sha256="x" * 64,
            schema_version="9",
        )


def test_content_plan_requires_mode_and_valid_selection() -> None:
    plan = ContentPlan(
        task_type="docx",
        mode="conservative",
        selections=[Selection(source_id="b0", purpose="正文", target="section/1")],
        exclusions=[],
    )
    assert plan.mode == "conservative"
    assert plan.excluded_ids() == set()
    with pytest.raises(ValidationError):
        ContentPlan(task_type="docx", mode="auto")  # 非法 mode


def test_content_plan_schema_version_gated() -> None:
    with pytest.raises(ValidationError):
        ContentPlan(task_type="docx", mode="vision", schema_version="2")


# ---------------- 隔离存储：同名 / 同 stem ----------------

def test_material_id_includes_full_hash_and_suffix(tmp_path: Path) -> None:
    a = _write(tmp_path / "a" / "report.md", "# 标题\n正文")
    b = _write(tmp_path / "b" / "report.md", "# 标题\n正文")
    # 内容相同 → 同 hash；不同内容 → 不同 hash
    assert material_id_for(a) == material_id_for(b)
    _write(tmp_path / "b" / "report.md", "# 改\n正文")
    assert material_id_for(a) != material_id_for(b)
    assert material_id_for(a).endswith(".md")
    assert len(material_id_for(a)) == 64 + len(".md")


def test_same_basename_different_dirs_do_not_overwrite(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    svc = MaterialService(ws)
    m1 = _write(tmp_path / "d1" / "相同.md", "# 文档一\n内容甲")
    m2 = _write(tmp_path / "d2" / "相同.md", "# 文档二\n内容乙")
    ir1 = svc.prepare_material(m1)
    ir2 = svc.prepare_material(m2)
    # 内容不同 → 不同 material_id → 不同 sources 子目录
    assert ir1.material_id != ir2.material_id
    source_dirs = [p.name for p in (ws / "sources").iterdir()]
    assert len(source_dirs) == 2
    # 两个原始名称都被记录
    names = {n for m in svc.catalog().manifest["materials"] for n in m["original_names"]}
    assert names == {"相同.md"}


def test_same_stem_different_extension_coexist(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    svc = MaterialService(ws)
    a = _write(tmp_path / "材料.md", "# 材料\nmd 内容")
    b = _write(tmp_path / "材料.txt", "txt 内容")
    ir_a = svc.prepare_material(a)
    ir_b = svc.prepare_material(b)
    assert ir_a.source_format == "md"
    assert ir_b.source_format == "txt"
    # 各自 material_id 目录互不覆盖：md 投影以首标题为题，txt 以文件名+正文
    md_dir = material_id_dir(ws, ir_a.material_id)
    txt_dir = material_id_dir(ws, ir_b.material_id)
    assert md_dir != txt_dir
    assert (md_dir / "content.md").read_text(encoding="utf-8").startswith("# 材料")
    assert "txt 内容" in (txt_dir / "content.md").read_text(encoding="utf-8")


# ---------------- Markdown 相对资源与 warning ----------------

def test_markdown_relative_images_materialized_with_location(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    svc = MaterialService(ws)
    src_dir = tmp_path / "src"
    _write(src_dir / "images" / "fig1.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 24)
    source = _write(src_dir / "doc.md", "# 报告\n\n正文段落。\n\n![示意图](images/fig1.png)")
    ir = svc.prepare_material(source)

    assert len(ir.assets) == 1
    asset = ir.assets[0]
    assert asset.sha256  # 内容 hash 非空
    # 图片块保留资产位置引用
    image_blocks = [b for b in ir.blocks if b.type == "image"]
    assert len(image_blocks) == 1
    assert image_blocks[0].asset_ids == [asset.id]
    # 资产路径是 workspace 相对路径且文件真实存在
    assert not Path(asset.path).is_absolute()
    assert (ws / asset.path).is_file()
    # 投影 content.md 引用该相对路径
    md = (material_id_dir(ws, ir.material_id) / "content.md").read_text(
        encoding="utf-8"
    )
    assert asset.path in md


def test_markdown_same_basename_assets_do_not_overwrite(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source_dir = tmp_path / "source"
    _write(source_dir / "left" / "figure.png", b"left-image")
    _write(source_dir / "right" / "figure.png", b"right-image")
    source = _write(
        source_dir / "report.md",
        "![Left](left/figure.png)\n\n![Right](right/figure.png)",
    )

    ir = MaterialService(workspace).prepare_material(source)

    assert len(ir.assets) == 2
    assert len({asset.id for asset in ir.assets}) == 2
    assert len({asset.path for asset in ir.assets}) == 2
    assert {asset.sha256 for asset in ir.assets} == {
        material_id_for(source_dir / "left" / "figure.png").split(".")[0],
        material_id_for(source_dir / "right" / "figure.png").split(".")[0],
    }


def test_markdown_missing_and_escaping_refs_warn(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    svc = MaterialService(ws)
    src_dir = tmp_path / "src"
    source = _write(
        src_dir / "doc.md",
        "# 报告\n\n![存在](ok.png)\n\n![缺失](missing.png)\n\n"
        "![越界](../secret.png)",
    )
    _write(src_dir / "ok.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 24)
    ir = svc.prepare_material(source)

    assert len(ir.assets) == 1  # 只物化存在的
    warns = "\n".join(ir.warnings)
    assert "缺失" in warns or "missing.png" in warns
    assert "逃出" in warns or "越界" in warns


def test_markdown_remote_and_data_urls_ignored(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    svc = MaterialService(ws)
    source = _write(
        tmp_path / "doc.md",
        "# 报告\n\n![网图](https://example.com/x.png)\n\n"
        "![data](data:image/png;base64,abc)",
    )
    ir = svc.prepare_material(source)
    assert ir.assets == []
    assert ir.blocks == [] or all(b.type != "image" for b in ir.blocks)


def test_markdown_table_formula_code_survive_projection(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    svc = MaterialService(ws)
    source = _write(
        tmp_path / "doc.md",
        "# 标题\n\n| 列A | 列B |\n| --- | --- |\n| 1 | 2 |\n\n"
        "公式 $E=mc^2$\n\n```python\nprint('hi')\n```",
    )
    ir = svc.prepare_material(source)
    types = [b.type for b in ir.blocks]
    assert "heading" in types
    assert "table" in types
    assert "code" in types
    table = next(b for b in ir.blocks if b.type == "table")
    assert "| 列A | 列B |" in table.text
    code = next(b for b in ir.blocks if b.type == "code")
    assert "print('hi')" in code.text
    # 代码块在兼容投影中保持围栏结构（语言标注不进 Block.level，int 字段）
    assert "```" in project_content_md(ir)
    assert "print('hi')" in project_content_md(ir)
    # 公式作为段落保留（不伪造结构）
    formula_para = next(b for b in ir.blocks if "$E=mc^2$" in b.text)
    assert formula_para.type == "paragraph"


# ---------------- 单次解析复用 ----------------

def test_same_hash_parsed_once(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    svc = MaterialService(ws)
    m1 = _write(tmp_path / "d1" / "报告.md", "# 标题\n正文")
    m2 = _write(tmp_path / "d2" / "副本.md", "# 标题\n正文")  # 同内容，不同文件名
    ir1 = svc.prepare_material(m1)
    ir2 = svc.prepare_material(m2)
    # 同 hash → 同 material_id → 返回同一缓存实例
    assert ir1.material_id == ir2.material_id
    assert ir1 is ir2
    # manifest 记录两个原始名称，但 sources 只有一个 hash 目录
    sources_dirs = [p.name for p in (ws / "sources").iterdir() if p.is_dir()]
    assert len(sources_dirs) == 1
    # hash 目录内只登记一份 original.md
    assert sorted(p.name for p in (ws / "sources" / sources_dirs[0]).iterdir()) == ["original.md"]
    entry = next(m for m in svc.catalog().manifest["materials"] if m["material_id"] == ir1.material_id)
    assert sorted(entry["original_names"]) == ["副本.md", "报告.md"]
    # 二次 prepare 同一文件仍命中缓存（不重复解析）
    ir3 = svc.prepare_material(m1)
    assert ir3 is ir1


def test_material_error_codes(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    svc = MaterialService(ws)
    # 不支持的扩展名（文件存在）→ MATERIAL_UNSUPPORTED
    bad = _write(tmp_path / "材料.xyz", "whatever")
    with pytest.raises(MaterialError) as exc:
        svc.prepare_material(bad)
    assert exc.value.code == "MATERIAL_UNSUPPORTED"
    # 不存在的文件 → MATERIAL_CORRUPT
    with pytest.raises(MaterialError) as exc:
        svc.prepare_material(tmp_path / "不存在.pdf")
    assert exc.value.code == "MATERIAL_CORRUPT"


def test_standalone_image_registers_asset(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    image = tmp_path / "portrait.png"
    image.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 24)

    ir = MaterialService(workspace).prepare_material(image)

    assert ir.source_format == "image"
    assert ir.blocks == []
    assert len(ir.assets) == 1
    assert ir.assets[0].mime_type == "image/png"
    assert not Path(ir.assets[0].path).is_absolute()


def test_inline_formula_keeps_surrounding_text(tmp_path: Path) -> None:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    source = tmp_path / "metrics.docx"
    document = Document()
    paragraph = document.add_paragraph("内部测试集中检索准确率达到约 ")
    paragraph._p.append(parse_xml(f'<m:oMath {nsdecls("m")}><m:r><m:t>93%</m:t></m:r></m:oMath>'))
    paragraph.add_run("，用于评估召回效果。")
    document.save(source)
    ir = MaterialService(tmp_path / "workspace").prepare_material(source)
    assert any(block.text == "内部测试集中检索准确率达到约 93%，用于评估召回效果。" and block.type == "paragraph" for block in ir.blocks)


def test_docx_ir_preserves_rich_block_order_and_image_links(tmp_path: Path) -> None:
    import base64

    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    image = tmp_path / "figure.png"
    image.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Z2S8AAAAASUVORK5CYII="
        )
    )
    source = tmp_path / "rich.docx"
    document = Document()
    document.add_heading("Heading", level=1)
    document.add_paragraph("First item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "1"
    document.add_paragraph().add_run().add_picture(str(image))
    formula = document.add_paragraph()
    formula._p.append(
        parse_xml(
            f'<m:oMath {nsdecls("m")}><m:r><m:t>E=mc^2</m:t></m:r></m:oMath>'
        )
    )
    document.save(source)

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    ir = MaterialService(workspace).prepare_material(source)

    types = [block.type for block in ir.blocks]
    assert "heading" in types
    assert "list" in types
    assert "table" in types
    assert "image" in types
    assert "formula" in types
    image_block = next(block for block in ir.blocks if block.type == "image")
    assert image_block.asset_ids == [ir.assets[0].id]
    assert ir.block_by_id(image_block.id) is image_block
    assert "| Name | Value |" in next(
        block.text for block in ir.blocks if block.type == "table"
    )


def test_pdf_preprocessing_exposes_single_prepared_docx(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import subprocess

    from docx import Document

    source = tmp_path / "report.pdf"
    source.write_bytes(b"%PDF-1.7\nfixture")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    # 注入隔离缓存根目录：跨任务持久缓存可能命中真实默认缓存（同 hash PDF），
    # 导致调用数被缓存吸收（实施计划 §6.4 允许缓存命中 0 调用）；测试必须
    # 保证"首次转换必然调用一次"的断言可复现。
    service = MaterialService(
        workspace,
        mineru_token="test-token",
        cache_root=tmp_path / "isolated-cache",
    )
    calls: list[list[str]] = []

    def fake_run(command: list[str], **_kwargs):  # type: ignore[no-untyped-def]
        calls.append(command)
        output = Path(command[3])
        document = Document()
        document.add_heading("Converted", level=1)
        document.save(output)
        return subprocess.CompletedProcess(command, 0, b"", b"")

    monkeypatch.setattr(service, "_run_subprocess", fake_run)

    ir = service.prepare_material(source)

    entry = service.catalog().manifest["materials"][0]
    assert ir.source_format == "pdf"
    assert len(calls) == 1
    assert Path(calls[0][1]).resolve() == Path("backend/skill_toolbox/parsers/mineru_pdf.py").resolve()
    assert calls[0][-1] == "--internal-absolute-paths"
    assert entry["prepared_docx"].endswith("report.docx")
    assert (workspace / entry["prepared_docx"]).is_file()

    # The relocated internal adapter preserves cache identity across workspaces.
    second_workspace = tmp_path / "second-workspace"
    second_workspace.mkdir()
    second = MaterialService(
        second_workspace, mineru_token="test-token",
        cache_root=tmp_path / "isolated-cache",
    )
    monkeypatch.setattr(second, "_run_subprocess", fake_run)
    cached_ir = second.prepare_material(source)
    assert cached_ir.source_format == "pdf"
    assert len(calls) == 1
    assert not (second_workspace / "artifacts").exists()


def test_material_subprocess_timeout_kills_child(tmp_path: Path) -> None:
    import subprocess
    import sys
    import time

    child = tmp_path / "late_writer.py"
    marker = tmp_path / "survived.txt"
    child.write_text(
        "import pathlib,sys,time\ntime.sleep(2)\npathlib.Path(sys.argv[1]).write_text('alive')\n",
        encoding="utf-8",
    )
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    service = MaterialService(workspace)

    with pytest.raises(subprocess.TimeoutExpired):
        service._run_subprocess(
            [sys.executable, str(child), str(marker)], timeout=0.05
        )
    time.sleep(2.2)

    assert not marker.exists()
    assert service._active_procs == []


# ---------------- 兼容投影单向生成 ----------------

def test_compat_projection_generated_from_ir(tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    svc = MaterialService(ws)
    source = _write(tmp_path / "doc.md", "# 标题\n\n正文段落")
    ir = svc.prepare_material(source)
    catalog = svc.catalog()
    md_rel = catalog.compat_projection[ir.material_id]
    md_path = ws / md_rel
    assert md_path.is_file()
    assert "# 标题" in md_path.read_text(encoding="utf-8")
    # 投影是 IR 的单向输出：修改投影不影响 IR
    md_path.write_text("# 被篡改\n", encoding="utf-8")
    assert ir.block_by_id(ir.blocks[0].id).text == "标题"


def test_same_bytes_different_formats_have_unique_ids_and_directories(
    tmp_path: Path,
) -> None:
    ws = tmp_path / "ws"
    ws.mkdir()
    service = MaterialService(ws)
    markdown = _write(tmp_path / "same.md", "same bytes")
    text = _write(tmp_path / "same.txt", "same bytes")

    md_ir = service.prepare_material(markdown)
    txt_ir = service.prepare_material(text)

    assert md_ir.material_id != txt_ir.material_id
    assert material_id_dir(ws, md_ir.material_id) != material_id_dir(
        ws, txt_ir.material_id
    )
    assert {block.id for block in md_ir.blocks}.isdisjoint(
        block.id for block in txt_ir.blocks
    )


def test_safe_stem_sanitizes() -> None:
    assert safe_stem("../报告 v1") == "报告_v1"
    assert safe_stem("a/b") == "a_b"
    # 分隔符/危险字符被替换，结果不含 / 与空字节
    assert "/" not in safe_stem("a/b")
    assert "\x00" not in safe_stem("x\x00y")
