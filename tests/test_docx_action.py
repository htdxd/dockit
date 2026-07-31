import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from skill_toolbox.actions.docx import build_simple_docx


def test_build_simple_docx_contains_required_elements(tmp_path: Path) -> None:
    spec_path = tmp_path / "spec.json"
    output_path = tmp_path / "simple-docx-test.docx"
    spec = {
        "title": "Skill Toolbox 文档能力测试",
        "summary": "用于验证基础 DOCX 产物。",
        "table": {
            "headers": ["能力", "状态"],
            "rows": [["表格", "通过"], ["公式", "通过"]],
        },
        "formula": "E = mc²",
        "code": 'def greet(name: str) -> str:\n    return f"Hello, {name}!"',
    }
    spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")

    build_simple_docx(spec_path, output_path)

    document = Document(output_path)
    assert document.paragraphs[0].text == "Skill Toolbox 文档能力测试"
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "表格"
    assert "def greet" in "\n".join(p.text for p in document.paragraphs)

    with ZipFile(output_path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "m:oMath" in xml
    assert "w:shd" in xml
