"""DOCX 主解析、显式备用视图和原件图片相互独立。"""
import hashlib
import json
import pytest
from docx import Document
from PIL import Image
from skill_toolbox.materials import MaterialService, MaterialError
from skill_toolbox.tools.materials import MaterialPlanService
from skill_toolbox.tools.mineru import MineruService
from skill_toolbox.contracts.common import ToolError


def test_docx_mineru_primary_and_native_backup_keep_original_images(tmp_path, monkeypatch):
    source = tmp_path / 'source.docx'
    original_image = tmp_path / 'original.png'
    generated_image = tmp_path / 'generated.png'
    Image.new('RGB', (80, 100), 'blue').save(original_image)
    Image.new('RGB', (40, 40), 'red').save(generated_image)
    doc = Document(); doc.add_paragraph('原件文字 93%'); doc.add_picture(str(original_image)); doc.save(source)
    calls = []
    def convert(self, staged, output, options):
        calls.append(staged)
        doc = Document(); doc.add_paragraph('MinerU 整理后的主内容'); doc.add_picture(str(generated_image)); doc.save(output)
        return {'status':'cache_miss', 'cache_key':'test'}
    monkeypatch.setattr(MineruService, 'convert', convert)
    workspace = tmp_path / 'workspace'
    material = MaterialService(workspace)
    ir = material.prepare_material(source)
    assert len(calls) == 1 and calls[0].suffix == '.docx'
    assert 'MinerU 整理后的主内容' in '\n'.join(b.text for b in ir.blocks)
    assert '原件文字' not in '\n'.join(b.text for b in ir.blocks)
    assert ir.source_format == 'docx'
    assert any('NATIVE_TEXT_RECOMMENDED' in warning for warning in ir.warnings)
    assert material.catalog().manifest['materials'][0]['primary_parser'] == 'mineru'
    saved = ir.model_dump_json()
    tools = MaterialPlanService(workspace, material.catalog(), 'resume', {'vision':True})
    assert tools.material_summary(ir.material_id).data['supplementary_views'] == ['native_text']
    backup = tools.read_material(ir.material_id, view='native_text', limit=1).data
    assert backup['blocks'][0]['text'] == '原件文字 93%'
    assert ir.model_dump_json() == saved and len(calls) == 1
    from types import SimpleNamespace
    from skill_toolbox.tools.resume_workflow import ResumeWorkflow
    from skill_toolbox.resume_task import operation_result
    workflow = ResumeWorkflow(SimpleNamespace(workspace=workspace), tools, 't001')
    result = tools.read_material(ir.material_id, view='native_text')
    workflow.record_native_read(result)
    output = operation_result('read_material', result)
    row = json.loads(output.content)['data']['blocks'][0]
    assert output.success and workflow.facts.sources[row['source_id']]['text'] == '原件文字 93%'
    original = next(a for a in ir.assets if a.source_locator.startswith('docx-original:'))
    assert (original.width, original.height) == (80,100)
    assert (workspace / original.path).read_bytes() == original_image.read_bytes()
    assert len(ir.assets) == 2
    for asset in ir.assets:
        assert hashlib.sha256((workspace / asset.path).read_bytes()).hexdigest() == asset.sha256


def test_docx_mineru_failure_is_not_silently_replaced_by_native_text(tmp_path, monkeypatch):
    source=tmp_path/'source.docx'; doc=Document(); doc.add_paragraph('原件'); doc.save(source)
    def fail(*args): raise ToolError('MINERU_PARSE_FAILED', 'conversion failed')
    monkeypatch.setattr(MineruService, 'convert', fail)
    with pytest.raises(MaterialError) as exc:
        MaterialService(tmp_path/'workspace').prepare_material(source)
    assert exc.value.code == 'MINERU_PARSE_FAILED'
